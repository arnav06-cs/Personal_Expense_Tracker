import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = create_element('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def update_report():
    doc_path = r'C:\Users\aaap7\OneDrive\Desktop\Arnav_VT_Report.docx'
    doc = docx.Document(doc_path)

    print("Total paragraphs in original doc:", len(doc.paragraphs))

    # Find the starting paragraph for Chapter I (index 124) and end before Student Copy (index 221)
    start_idx = -1
    end_idx = -1

    for i, p in enumerate(doc.paragraphs):
        if "Chapter I:" in p.text or "CHAPTER I" in p.text or "Chapter I : INTRODUCTION" in p.text:
            start_idx = i
            break

    for i, p in enumerate(doc.paragraphs):
        if "Acknowledgement of Vocational Training" in p.text:
            end_idx = i
            break

    print(f"Replacing paragraphs from index {start_idx} to {end_idx}")

    # Clear existing chapter paragraphs between start_idx and end_idx
    # In python-docx, deleting paragraphs from back to front is safe
    for i in range(end_idx - 1, start_idx - 1, -1):
        p = doc.paragraphs[i]._element
        p.getparent().remove(p)

    # Now we insert our updated content at start_idx position
    # To insert before the student copy, we can insert before doc.paragraphs[start_idx]
    ref_p = doc.paragraphs[start_idx]

    def add_p(text, style='Normal', space_before=6, space_after=6, bold=False, italic=False, font_size=12, color=RGBColor(0,0,0), align=WD_ALIGN_PARAGRAPH.LEFT):
        p = ref_p.insert_paragraph_before()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(font_size)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = color
        return p

    def add_h1(text):
        return add_p(text, space_before=18, space_after=12, bold=True, font_size=16, color=RGBColor(15, 23, 42))

    def add_h2(text):
        return add_p(text, space_before=14, space_after=8, bold=True, font_size=14, color=RGBColor(16, 185, 129))

    def add_h3(text):
        return add_p(text, space_before=10, space_after=6, bold=True, font_size=12, color=RGBColor(30, 41, 59))

    def add_bullet(text, bold_prefix=""):
        p = ref_p.insert_paragraph_before()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        
        run_b = p.add_run("•  " + bold_prefix + " ") if bold_prefix else p.add_run("•  ")
        run_b.font.name = 'Calibri'
        run_b.font.size = Pt(11)
        run_b.bold = True if bold_prefix else False
        run_b.font.color.rgb = RGBColor(15, 23, 42)

        run_t = p.add_run(text)
        run_t.font.name = 'Calibri'
        run_t.font.size = Pt(11)
        run_t.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def add_body(text):
        return add_p(text, space_before=4, space_after=6, font_size=11, color=RGBColor(51, 65, 85))

    # ==========================================
    # CHAPTER I: INTRODUCTION
    # ==========================================
    add_h1("Chapter I: INTRODUCTION")

    add_h2("1.1 About the Organization – Infynas Learning Solutions")
    add_body(
        "Infynas Learning Solutions, Raipur is a leading technology training and skill-development organization "
        "that provides hands-on, industry-oriented training in software engineering to Computer Science and "
        "Information Technology students. The organization focuses on bridging the gap between theoretical academic "
        "curricula and modern industrial software practices. Trainees are exposed to professional web engineering workflows, "
        "Model-View-Template (MVT) design patterns, version control, database design, and end-to-end full-stack web application development."
    )
    add_body(
        "During the vocational training period, candidates work on production-grade projects under the direct supervision of "
        "experienced industry mentors. The training emphasizes writing modular, maintainable code, adhering to Web Content "
        "Accessibility Guidelines (WCAG), implementing responsive user interfaces, and ensuring robust backend security."
    )

    add_h2("1.2 Objective and Scope of the Project")
    add_body(
        "The primary objective of the Vocational Training project was to architect, build, and evaluate NovaSpend – "
        "an AI-powered, full-stack Personal Expense Tracker & Financial Command Center. Traditional manual methods of expense "
        "logging—such as physical paper diaries or basic offline spreadsheets—are error-prone, time-consuming, and lack automated "
        "visual analytics, real-time budget warnings, or intelligent spending velocity insights."
    )
    add_body("The specific functional and technical objectives of the NovaSpend project include:")
    
    add_bullet("Provides a secure dark glassmorphic user authentication system for account registration, sign-in, and session management.", "User Authentication & Portal:")
    add_bullet("Enables users to log daily expenses with title, amount, category, payment method (UPI, Cash, Credit Card, Bank Transfer), date, notes, and optional physical receipt image attachments.", "Itemized Expense Tracking:")
    add_bullet("Includes a dedicated module to record income sources (Salary, Freelancing, Investments) and dynamically compute net monthly balance.", "Itemized Income Logging:")
    add_bullet("Allows users to create, update, and organize custom categories with hex color coding and confirmation modals.", "Custom Category Management:")
    add_bullet("Features interactive KPI cards, monthly/yearly totals, and a Chart.js doughnut chart for real-time category spending distribution.", "Executive Dashboard Analytics:")
    add_bullet("Integrates a real-time LLM-powered financial advisor that computes budget velocity, performs affordability checks (e.g., 'Can I afford a ₹2,000 purchase?'), and delivers instant spending diagnostics.", "NovaAI Financial Advisor:")
    add_bullet("Incorporates 5 Saver Ranks (Level 1 Financial Explorer to Level 5 Financial Titan), Experience Points (XP) level progression, 10 unlockable achievement badges, saver streaks, and weekly challenges with an interactive Level Roadmap modal.", "Savings Arcade & Gamification:")
    add_bullet("Supports global currency display toggling (INR ₹, USD $, EUR €, GBP £, AED) and an interactive real-time exchange rates calculation panel with FX matrix tables.", "Multi-Currency & FX Engine:")
    add_bullet("Calculates per-person split amounts for shared restaurant bills and group expenses with one-click category logging.", "Group Bill Splitter Tool:")
    add_bullet("Generates executive PDF reports via ReportLab (with summary KPI boxes and zebra-striped ledgers) and Microsoft Excel workbooks via OpenPyXL (with currency number formatting).", "Executive PDF & Excel Reports:")
    add_bullet("Incorporate a custom dark glassmorphic admin site header and registered models for system administration.", "NovaSpend Admin Console:")

    add_body(
        "The scope of the project encompasses a multi-user secure web platform where each user's financial records, "
        "categories, budgets, and arcade achievements are strictly isolated via Django's Object-Relational Mapping (ORM) "
        "and request-level authorization decorators."
    )

    add_h2("1.3 Tools and Technologies Used")
    add_body("The project was developed using state-of-the-art open-source software technologies:")
    add_bullet("Python 3.12 & Django 5.1 Web Framework", "Backend Server:")
    add_bullet("SQLite3 Relational Database Engine (Django ORM)", "Database:")
    add_bullet("HTML5, Vanilla CSS3 (Custom Dark Glassmorphic Design System), JavaScript (ES6+ Async Fetch)", "Frontend:")
    add_bullet("Bootstrap 5.3 Framework & Bootstrap Icons 1.11", "UI Components:")
    add_bullet("Chart.js 4.4 JavaScript Charting Engine", "Data Visualization:")
    add_bullet("ReportLab (PDF Generation Toolkit) & OpenPyXL (Spreadsheet Engine)", "Export Engines:")

    # ==========================================
    # CHAPTER II: HARDWARE AND SOFTWARE REQUIREMENTS
    # ==========================================
    add_h1("Chapter II: HARDWARE AND SOFTWARE REQUIREMENTS")

    add_h2("2.1 Software Requirements")
    add_body("The development and operational deployment environment requires the following software components:")
    add_bullet("Windows 10 / Windows 11 (64-bit) or Linux / macOS", "Operating System:")
    add_bullet("Python 3.12.x Runtime Environment", "Programming Language:")
    add_bullet("Django Framework 5.1.x", "Web Framework:")
    add_bullet("ReportLab 4.x, OpenPyXL 3.x, Widget-Tweaks, Pillow", "Python Dependencies:")
    add_bullet("SQLite3 Relational Database", "Database Management:")
    add_bullet("Google Chrome, Microsoft Edge, or Mozilla Firefox (ES6 compatible)", "Client Web Browser:")
    add_bullet("Visual Studio Code / Antigravity IDE", "Integrated Development Environment:")

    add_h2("2.2 Hardware Requirements")
    add_body("The minimal hardware configurations required to execute and serve the web application locally are:")
    add_bullet("Intel Core i3 / AMD Ryzen 3 (Dual-Core 2.0 GHz) or higher", "Processor:")
    add_bullet("4 GB minimum (8 GB recommended for concurrent development)", "System Memory (RAM):")
    add_bullet("500 MB free hard disk space for project source code, database, and receipt uploads", "Storage Space:")
    add_bullet("1366 x 768 pixels or higher (Responsive layout adapts seamlessly to mobile devices)", "Display Resolution:")

    # ==========================================
    # CHAPTER III: METHODOLOGY & DATABASE DESIGN
    # ==========================================
    add_h1("Chapter III: FLOW CHART / E-R DIAGRAMS / METHODOLOGY")

    add_h2("3.1 Methodology & Application Architecture")
    add_body(
        "NovaSpend is architected around Django's Model-View-Template (MVT) software engineering pattern. "
        "The architecture enforces separation of concerns between database models, business logic views, and presentation templates."
    )
    add_body("The operational execution pipeline follows these steps:")
    add_bullet("Users register or authenticate via Django's secure PBKDF2 password hashing system.", "1. Authentication:")
    add_bullet("Upon login, Django inspects session state and routes the user to the Executive Dashboard.", "2. Dispatching:")
    add_bullet("Views execute database queries via Django ORM scoped to request.user to compute totals, budget limits, and category aggregations.", "3. View Execution:")
    add_bullet("The computed context is rendered through dark glassmorphic HTML templates enhanced with Bootstrap 5.3 and Chart.js.", "4. UI Rendering:")
    add_bullet("Asynchronous JavaScript (Fetch API) handles real-time NovaAI Advisor queries and conversion rate panel calculations without page reloads.", "5. Async Interactivity:")

    add_h2("3.2 Database Design & E-R Model")
    add_body("The system database consists of six relational models designed with foreign-key constraints and index optimizations:")
    add_bullet("Django built-in model handling user credentials and permissions.", "User Model:")
    add_bullet("OneToOne link to User storing preferred currency symbol, avatar image, and default budget.", "UserProfile Model:")
    add_bullet("ForeignKey link to User storing category names and hex color codes.", "Category Model:")
    add_bullet("ForeignKey link to User & Category storing title, amount, date, payment method, receipt image, and notes.", "Expense Model:")
    add_bullet("ForeignKey link to User storing title, amount, income date, source, and notes.", "Income Model:")
    add_bullet("ForeignKey link to User storing month, year, and target monthly spending budget amount.", "Budget Model:")

    # ==========================================
    # CHAPTER IV: RESULTS & DISCUSSIONS
    # ==========================================
    add_h1("Chapter IV: RESULTS & DISCUSSIONS")

    add_h2("4.1 Module-wise Implementation")
    add_body("The NovaSpend Personal Expense Tracker was built and verified across 11 integrated modules:")
    add_bullet("Dark glassmorphic portal for registration, login, logout, and password management.", "1. User Authentication Module:")
    add_bullet("Central hub presenting spending KPIs, budget progress meter, and Chart.js doughnut visual charts.", "2. Executive Dashboard Module:")
    add_bullet("Itemized expense creation, updating, deleting, search filtering, and receipt image attachments.", "3. Expense Management Module:")
    add_bullet("Record and monitor income entries (Salary, Freelance, Bonus) with net monthly savings computation.", "4. Income Logging Module:")
    add_bullet("Custom category creation with color pickers, edit modals, and safety confirmation popups.", "5. Category Management Module:")
    add_bullet("Configure monthly spending targets, track utilization bars, and trigger alert warnings.", "6. Monthly Budget Command Center:")
    add_bullet("Real-time AI assistant for budget analysis, affordability checks, and spending health diagnostics.", "7. NovaAI Financial Advisor:")
    add_bullet("XP level progression (Level 1 to 5), 10 unlockable badges, saver streaks, and interactive Level Roadmap modal.", "8. Savings Arcade & Gamification:")
    add_bullet("Switch active display currency (INR ₹, USD $, EUR €, GBP £, AED) and calculate live FX exchange rates.", "9. Multi-Currency & Conversion Rate Panel:")
    add_bullet("Split shared restaurant and group expenses among friends with one-click category logging.", "10. Group Bill Splitter Tool:")
    add_bullet("Export filtered expense ledger into executive PDF documents (ReportLab) or Excel workbooks (OpenPyXL).", "11. PDF & Excel Export Engine:")

    add_h2("4.2 Verification and System Outputs")
    add_body(
        "System integration testing was executed locally using Django's WSGI server. All 11 modules passed verification with zero errors. "
        "Database query integrity was validated under multi-user testing, confirming strict isolation of user data. Exporter modules "
        "were verified by rendering executive PDF reports and formatted Excel sheets with currency formatting."
    )

    # ==========================================
    # CHAPTER V: CONCLUSION & FUTURE SCOPE
    # ==========================================
    add_h1("Chapter V: CONCLUSION & SCOPE OF FURTHER WORK")

    add_h2("5.1 Conclusion")
    add_body(
        "The Vocational Training at Infynas Learning Solutions, Raipur provided invaluable practical experience in full-stack web application "
        "engineering. The developed NovaSpend Personal Expense Tracker successfully solves real-world personal finance management challenges by "
        "combining modern dark glassmorphic design, dynamic database tracking, gamified behavioral incentives, multi-currency support, and AI financial advice."
    )

    add_h2("5.2 Future Scope")
    add_body("Future enhancements planned for the NovaSpend platform include:")
    add_bullet("Automated OCR extraction from uploaded receipt images using Optical Character Recognition.", "1. Automated Receipt OCR:")
    add_bullet("Direct API integration with banking gateways for automated transaction syncing.", "2. Bank Statement Integration:")
    add_bullet("Automated tracking and cancellation alerts for recurring software subscriptions.", "3. Recurring Subscriptions Tracker:")
    add_bullet("Exposing DRF REST endpoints to build native Android and iOS mobile companion apps.", "4. Mobile Native App (Flutter):")
    add_bullet("Cloud deployment to Google Cloud Platform (GCP) or Render with PostgreSQL production database.", "5. Cloud Production Deployment:")

    # ==========================================
    # REFERENCES
    # ==========================================
    add_h1("REFERENCES")
    add_bullet("Django Software Foundation, 'Django 5.1 Documentation', https://docs.djangoproject.com/", "1.")
    add_bullet("Python Software Foundation, 'Python 3.12 Documentation', https://docs.python.org/3/", "2.")
    add_bullet("Bootstrap Team, 'Bootstrap 5.3 Documentation & Components', https://getbootstrap.com/", "3.")
    add_bullet("Chart.js Contributors, 'Chart.js v4.4 HTML5 Canvas Charting Documentation', https://www.chartjs.org/", "4.")
    add_bullet("ReportLab Inc., 'ReportLab PDF Toolkit User Guide', https://www.reportlab.com/", "5.")
    add_bullet("OpenPyXL Team, 'OpenPyXL - Python Library to Read/Write Excel 2010 xlsx Files', https://openpyxl.readthedocs.io/", "6.")
    add_bullet("MDN Web Docs, 'JavaScript ES6 & Modern Web APIs', https://developer.mozilla.org/", "7.")

    output_path = r'C:\Users\aaap7\OneDrive\Desktop\Arnav_VT_Report_Updated.docx'
    doc.save(output_path)
    print("Report successfully updated and saved to:", output_path)

    # Try copying to original file if Word is closed
    import shutil
    try:
        shutil.copy(output_path, doc_path)
        print("Also updated original file:", doc_path)
    except Exception as e:
        print("Note: Original file was open in Word, so saved to Arnav_VT_Report_Updated.docx!")

if __name__ == '__main__':
    update_report()
